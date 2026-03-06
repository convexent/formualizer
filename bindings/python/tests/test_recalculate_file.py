from pathlib import Path

import pytest

import formualizer as fz

try:
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover - allow skipping if not present in dev env
    openpyxl = None

pytestmark = pytest.mark.skipif(openpyxl is None, reason="openpyxl not installed")


def test_recalculate_file_in_place_summary_and_formula_preservation(tmp_path: Path):
    path = tmp_path / "recalc_in_place.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["B1"] = "=SUM(A1:A2)"
    ws["B2"] = "=_xlfn._xlws.SUM(A1:A2)"
    ws["B3"] = "=1/0"
    wb.save(path)

    summary = fz.recalculate_file(str(path))

    assert summary["status"] == "errors_found"
    assert summary["evaluated"] == 3
    assert summary["errors"] == 1
    assert summary["total_formulas"] == 3
    assert summary["total_errors"] == 1
    assert summary["sheets"]["Sheet1"] == {"evaluated": 3, "errors": 1}
    assert summary["error_summary"]["#DIV/0!"]["count"] == 1
    assert summary["error_summary"]["#DIV/0!"]["locations"] == ["Sheet1!B3"]

    data_only = openpyxl.load_workbook(path, data_only=True)
    with_formula = openpyxl.load_workbook(path, data_only=False)

    # Numeric cached values may be returned as numbers or numeric strings
    # depending on the active umya implementation.
    assert data_only["Sheet1"]["B1"].value in (3, 3.0, "3")
    assert data_only["Sheet1"]["B2"].value in (3, 3.0, "3")
    assert data_only["Sheet1"]["B3"].value == "#DIV/0!"

    # Formula text remains intact.
    assert with_formula["Sheet1"]["B1"].value == "=SUM(A1:A2)"
    assert with_formula["Sheet1"]["B2"].value == "=_xlfn._xlws.SUM(A1:A2)"
    assert with_formula["Sheet1"]["B3"].value == "=1/0"


def test_recalculate_file_output_path_keeps_input_unmodified(tmp_path: Path):
    in_path = tmp_path / "recalc_input.xlsx"
    out_path = tmp_path / "recalc_output.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 4
    ws["A2"] = 6
    ws["B1"] = "=A1+A2"
    wb.save(in_path)

    before = openpyxl.load_workbook(in_path, data_only=True)
    assert before["Sheet1"]["B1"].value is None

    summary = fz.recalculate_file(str(in_path), output=str(out_path))
    assert summary["status"] == "success"
    assert summary["evaluated"] == 1
    assert summary["errors"] == 0

    in_after = openpyxl.load_workbook(in_path, data_only=True)
    out_after = openpyxl.load_workbook(out_path, data_only=True)

    assert in_after["Sheet1"]["B1"].value is None
    assert out_after["Sheet1"]["B1"].value in (10, 10.0, "10")
